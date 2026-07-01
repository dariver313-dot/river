import tkinter as tk
from tkinter import filedialog, messagebox
import pandas as pd
from openpyxl import load_workbook
import os

class ReconciliationApp:
    def __init__(self, root):
        self.root = root
        self.root.title("澳门娱乐城 - 自动做账系统 v2.1")
        self.root.geometry("800x500")

        self.charge_path = tk.StringVar()
        self.withdraw_path = tk.StringVar()
        self.template_path = tk.StringVar()

        self.create_widgets()

    def create_widgets(self):
        # --- 文件选择区 ---
        tk.Label(self.root, text="1. 选择充值账变明细:").grid(row=0, column=0, padx=10, pady=10, sticky="w")
        tk.Entry(self.root, textvariable=self.charge_path, width=50).grid(row=0, column=1, padx=10, pady=10)
        tk.Button(self.root, text="浏览...", command=lambda: self.browse_file(self.charge_path)).grid(row=0, column=2, padx=10, pady=10)

        tk.Label(self.root, text="2. 选择提现订单明细:").grid(row=1, column=0, padx=10, pady=10, sticky="w")
        tk.Entry(self.root, textvariable=self.withdraw_path, width=50).grid(row=1, column=1, padx=10, pady=10)
        tk.Button(self.root, text="浏览...", command=lambda: self.browse_file(self.withdraw_path)).grid(row=1, column=2, padx=10, pady=10)

        tk.Label(self.root, text="3. 选择模版表 (澳门娱乐城.xlsx):").grid(row=2, column=0, padx=10, pady=10, sticky="w")
        tk.Entry(self.root, textvariable=self.template_path, width=50).grid(row=2, column=1, padx=10, pady=10)
        tk.Button(self.root, text="浏览...", command=lambda: self.browse_file(self.template_path)).grid(row=2, column=2, padx=10, pady=10)

        # --- 执行按钮 ---
        tk.Button(self.root, text="⚙️ 一键生成做账表", bg="#4CAF50", fg="white", font=("Arial", 14, "bold"),
                  command=self.process_data).grid(row=3, column=0, columnspan=3, pady=20)

        # --- 结果输出框 ---
        self.result_text = tk.Text(self.root, height=10, width=100, bg="#f0f0f0")
        self.result_text.grid(row=4, column=0, columnspan=3, padx=10, pady=5)

    def browse_file(self, path_var):
        filename = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx")])
        if filename:
            path_var.set(filename)

    def process_data(self):
        charge_file = self.charge_path.get()
        withdraw_file = self.withdraw_path.get()
        template_file = self.template_path.get()

        if not charge_file or not withdraw_file or not template_file:
            messagebox.showerror("错误", "请先选择充值、提现明细和模版表！")
            return

        try:
            self.result_text.delete(1.0, tk.END)
            self.result_text.insert(tk.END, "正在处理数据，请稍候...\n\n")
            self.root.update()

            # 1. 读取数据
            df_charge = pd.read_excel(charge_file)
            df_withdraw = pd.read_excel(withdraw_file)
            
            # 2. 加载模版
            wb = load_workbook(template_file)
            ws = wb['总汇']

            # ================== 提现对账逻辑 ==================
            success_withdraw = df_withdraw[df_withdraw['订单状态'] == '代付成功']
            
            withdraw_channel_map = {
                'AB': 'AB钱包', 'CB': 'C币钱包', 'KD': 'K豆钱包', 
                'JD': 'JD钱包', 'OK': 'OK钱包', '234': '234钱包',
                '988': '988钱包', '365': '365钱包'
            }
            
            withdraw_result = {}
            total_withdraw = 0

            for _, row in success_withdraw.iterrows():
                channel_name = str(row['操作说明'])
                amount = row['申请金额']
                total_withdraw += amount
                
                matched_account = None
                for key, value in withdraw_channel_map.items():
                    if key in channel_name:
                        matched_account = value
                        break
                
                if matched_account:
                    withdraw_result[matched_account] = withdraw_result.get(matched_account, 0) + amount

            # 写入提现数据到模版 (找到对应的帐户名行，写入J列-会员取款)
            for row_idx in range(4, 35): 
                account_name = ws[f'B{row_idx}'].value
                if account_name in withdraw_result:
                    ws[f'J{row_idx}'] = withdraw_result[account_name]

            # 写入提现总计
            ws['J53'] = total_withdraw

            # ================== 充值对账逻辑 ==================
            total_manual = 0 # 初始化人工充值总额

            # 1. 线上充值
            online_charge = df_charge[df_charge['变动类型'] == '线上充值转入']['变动金额'].sum()
            ws['F52'] = online_charge 

            # 2. 人工/群码充值 (这里修复了之前的引号错误)
            manual_charge = df_charge[df_charge['变动类型'] == '手动微信扫码']
            if not manual_charge.empty:
                grouped_manual = manual_charge.groupby('操作员备注')['变动金额'].sum()
                
                charge_account_map = {
                    '如意': '如意群码', '顺利': '顺利群码', '村长': '村长群码', '发财': '发财群码'
                }
                
                for name, amount in grouped_manual.items():
                    matched_account = charge_account_map.get(str(name))
                    if matched_account:
                        for row_idx in range(4, 35):
                            if ws[f'B{row_idx}'].value == matched_account:
                                ws[f'H{row_idx}'] = amount
                                break
                    total_manual += amount
                
                ws['H53'] = total_manual

            # ================== 保存文件 ==================
            save_dir = os.path.dirname(template_file)
            save_path = os.path.join(save_dir, "澳门娱乐城_已做账.xlsx")
            wb.save(save_path)

            result_msg = f"✅ 做账完成！\n\n"
            result_msg += f"📝 写入提款: 共计 {len(withdraw_result)} 个账户，总金额 {total_withdraw:,.2f}\n"
            result_msg += f"📝 写入充值: 线上 {online_charge:,.2f}, 人工 {total_manual:,.2f}\n\n"
            result_msg += f"💾 文件已保存至: \n{save_path}"
            
            self.result_text.delete(1.0, tk.END)
            self.result_text.insert(tk.END, result_msg)

        except Exception as e:
            messagebox.showerror("处理错误", f"文件处理失败，请检查格式：\n{str(e)}")

if __name__ == "__main__":
    root = tk.Tk()
    app = ReconciliationApp(root)
    root.mainloop()