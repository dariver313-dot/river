"""Excel 模版写入器 v2 — 基于真实总汇结构。

列布局: F=会员充值 G=未提交 H=人工充值 I=资金转出 J=会员取款 K=手续费 L=余额 M=昨日余额
合计行: Row 3
"""

import logging
import os
from datetime import date
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)


class ExcelWriter:
    """将对账结果写入平台模版 Excel。"""

    def __init__(self, platform_config: dict, output_config: dict = None):
        self.config = platform_config
        self.platform_name = platform_config["name"]
        self.template_path = platform_config["template"]
        self.cells = platform_config["template_cells"]

        output_config = output_config or {}
        self.output_dir = output_config.get("dir", "data/output")
        self.filename_pattern = output_config.get(
            "filename_pattern", "{platform_name}_已做账_{date}.xlsx"
        )

    def write(
        self,
        accounts: dict,
        totals: dict,
        target_date: Optional[str] = None,
        work_file: Optional[str] = None,
    ) -> str:
        """写入对账结果并保存。

        Args:
            accounts: {账户名: AccountRow}
            totals: 合计行数据
            target_date: 日期 YYYYMMDD
            work_file: 含后台明细的工作文件，None则用模版

        Returns:
            输出文件路径
        """
        if target_date is None:
            target_date = date.today().strftime("%Y%m%d")

        # 用工作文件(含后台明细)或模版
        source = work_file if work_file and os.path.exists(work_file) else self.template_path
        if not os.path.exists(source):
            raise FileNotFoundError(f"文件不存在: {source}")

        wb = load_workbook(source)
        ws = wb[self.cells["sheet"]]
        logger.info(f"写入总汇: {source}")

        # ★ 模板有公式，只写 M列(昨日余额) 和 N列(USDT昨日余额)
        # F-L列 = 公式自动引用 → 不覆盖
        # M列(昨日余额) = 模板无公式，需手动填入
        # N列(备注) = USDT区域复用为USDT昨日余额
        for r in range(self.cells["data_start_row"], self.cells["data_end_row"] + 1):
            # 普通账户(Row5-55): C列=帐户名
            account_name = ws[f"{self.cells['col_account']}{r}"].value
            if not account_name:
                # USDT区域(Row64-69): B列=编号(如USDT5)
                account_name = ws[f"B{r}"].value
            if not account_name:
                continue
            account_name = str(account_name).strip()
            # 跳过公式、标题、非账户行
            if account_name.startswith("=") or account_name.startswith("SUM"):
                continue

            row_data = accounts.get(account_name)
            if not row_data:
                continue

            # M列(人民币昨日余额)
            self._write_cell(ws, r, self.cells["col_yesterday"], row_data.yesterday_balance)
            # N列(USDT昨日余额) — USDT区域用N列存USDT期初
            if row_data.usdt_yesterday:
                self._write_cell(ws, r, self.cells.get("col_remark", "N"), row_data.usdt_yesterday)

        # ★ 合计行不写 — Row3 有 SUM 公式，自动计算

        # 3. 保存 (文件名如 2026天游6.3.xlsx)
        os.makedirs(self.output_dir, exist_ok=True)
        # date_short: 20260603 → 6.3
        if len(target_date) == 8:
            m = str(int(target_date[4:6]))
            d = str(int(target_date[6:8]))
            date_short = f"{m}.{d}"
        else:
            date_short = target_date
        filename = self.filename_pattern.format(
            platform_name=self.platform_name, date=target_date,
            date_short=date_short,
        )
        output_path = os.path.join(self.output_dir, filename)
        wb.save(output_path)

        n_accounts = len([a for a in accounts.values() if a.balance != 0 or a.yesterday_balance != 0])
        logger.info(
            f"{self.platform_name} 写入完成: {n_accounts} 个非零账户 | "
            f"余额总计 {totals.get('balance', 0):,.2f}"
        )
        logger.info(f"输出文件: {output_path}")

        return output_path

    def _write_cell(self, ws, row: int, col_letter: str, value: float):
        """写入单元格。表头行(row≤11)遇合并单元格跳过，不破坏模板结构。"""
        try:
            ws[f"{col_letter}{row}"] = value
        except AttributeError:
            # MergedCell: 表头区域不破坏，数据行尝试拆开
            if row <= 11:
                return
            for mr in list(ws.merged_cells.ranges):
                if mr.min_row <= row <= mr.max_row:
                    from openpyxl.utils import get_column_letter
                    if get_column_letter(mr.min_col) <= col_letter <= get_column_letter(mr.max_col):
                        ws.unmerge_cells(str(mr))
                        break
            try:
                ws[f"{col_letter}{row}"] = value
            except AttributeError:
                pass
